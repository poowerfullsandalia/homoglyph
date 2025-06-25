#!/usr/bin/env python3
"""
Batch Word (.docx) homoglyph transformer.

This utility scans the *input_docs* directory for Word documents, applies the
homoglyph substitution (using the core HomoglyphTransformer defined in
``homoglyph_poc.py``), and writes the modified files to *output_docs* while
preserving all Word layout/style information (paragraphs, runs, tables, etc.).

Only the textual content of each *run* is altered – all other formatting
properties (bold, italics, headings, lists, tables) remain untouched because
python-docx leaves run styling intact when we replace ``run.text``.

Usage (from repository root):

    python homoglyph_docx.py            # default 30 % replacement rate
    python homoglyph_docx.py -r 0.5     # 50 % replacement rate

The script creates *input_docs* and *output_docs* directories automatically if
they do not exist. Populate *input_docs* with .docx files to transform.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Tuple

try:
    from docx import Document  # type: ignore
except ModuleNotFoundError as exc:
    sys.stderr.write("python-docx is required: pip install python-docx\n")
    raise exc

# Import the existing homoglyph transformer without triggering its demo code.
from homoglyph_poc import HomoglyphTransformer

# Constants – change cautiously to respect coding-guideline style rules
INPUT_DIR = Path("input_docs")
OUTPUT_DIR = Path("output_docs")
MAX_FILES = 1000  # Simple fixed upper-bound to satisfy rule #2 (loop bounds)


def transform_docx(
    src_path: Path,
    dst_path: Path,
    transformer: HomoglyphTransformer,
    replacement_rate: float,
) -> int:
    """Transform *src_path* Word document and save to *dst_path*.

    Returns the number of character replacements performed.
    """
    document = Document(src_path)
    replacements = 0

    # Helper to transform all runs inside a paragraph-like container
    def _process_runs(runs):
        nonlocal replacements
        for run in runs:
            new_text, count = transformer.transform(run.text, replacement_rate)
            run.text = new_text
            replacements += count

    # Process paragraphs
    for para in document.paragraphs[: 10_000]:  # enforce static upper-bound
        _process_runs(para.runs)

    # Process tables (rows -> cells -> paragraphs)
    for tbl in document.tables[: 1_000]:
        for row in tbl.rows[: 1_000]:
            for cell in row.cells[: 1_000]:
                for para in cell.paragraphs[: 10_000]:
                    _process_runs(para.runs)

    # Save the modified document
    document.save(dst_path)
    return replacements


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Transform .docx files using homoglyph substitution.")
    parser.add_argument(
        "-r",
        "--rate",
        type=float,
        default=0.3,
        help="Replacement rate between 0 and 1 (default: 0.3)",
    )

    args = parser.parse_args(argv)
    rate = min(max(args.rate, 0.0), 1.0)  # Clamp to [0, 1]

    # Ensure I/O directories exist
    INPUT_DIR.mkdir(exist_ok=True)
    OUTPUT_DIR.mkdir(exist_ok=True)

    transformer = HomoglyphTransformer()

    docx_files = sorted(INPUT_DIR.glob("*.docx"))[:MAX_FILES]
    if not docx_files:
        print("[INFO] No .docx files found in", INPUT_DIR)
        print("       Populate the folder with Word documents and rerun.")
        return

    for idx, src in enumerate(docx_files, start=1):
        dst = OUTPUT_DIR / src.name
        count = transform_docx(src, dst, transformer, rate)
        print(f"[{idx}/{len(docx_files)}] {src.name}: {count} characters replaced → {dst}")

    print("\nDone. Transformed documents are available in", OUTPUT_DIR)


if __name__ == "__main__":
    main() 